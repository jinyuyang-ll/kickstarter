from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("Kickstarter爬虫HTTP403排查操作步骤与后续风险预测手册.docx")

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_YELLOW = "FFF4CE"
PALE_RED = "FCE8E6"
PALE_GREEN = "E6F4EA"
GRAY = "666666"
WHITE = "FFFFFF"
BLACK = "202124"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=11, bold=False, color=BLACK, font="Microsoft YaHei"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run(run, size=9, color=GRAY)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 27, NAVY, 0, 8),
        ("Subtitle", 13, DARK_BLUE, 0, 14),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_header_footer(doc):
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(hp.add_run("Kickstarter 爬虫运行与故障处理手册"), size=8.5, bold=True, color=GRAY)
    p_pr = hp._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D7DBE2")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    add_page_number(section.footer.paragraphs[0])


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(f"{label}："), bold=True, color=NAVY)
    set_run(p.add_run(text), color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_run(p.add_run(text))
    return p


def add_step(doc, title, detail):
    p = doc.add_paragraph(style="List Number")
    set_run(p.add_run(title), bold=True, color=NAVY)
    set_run(p.add_run(f"：{detail}"))
    return p


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F8FA")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        set_run(p.add_run(line), size=9, color="333333", font="Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    repeat_header(hdr)
    for idx, text in enumerate(headers):
        set_cell_shading(hdr.cells[idx], header_fill)
        p = hdr.cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(text), size=9.5, bold=True, color=NAVY)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(str(text)), size=9.2)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def build():
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)

    for _ in range(3):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(kicker.add_run("运行人员故障处理 SOP"), size=11, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(title.add_run("Kickstarter 爬虫 HTTP 403"), size=27, bold=True, color=NAVY)
    title2 = doc.add_paragraph(style="Title")
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(title2.add_run("排查操作步骤与后续风险预测手册"), size=24, bold=True, color=NAVY)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(
        subtitle.add_run("适用于“浏览器可访问，但 smoke_test.py 全部返回 security_challenge”的现场"),
        size=12.5,
        color=DARK_BLUE,
    )
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(meta.add_run(f"编制日期：{date.today().isoformat()}　适用环境：Windows 10/11、Python 3.10+"), size=9.5, color=GRAY)
    doc.add_paragraph()
    add_callout(
        doc,
        "当前判断",
        "程序并非卡死。日志表明 Kickstarter 在项目主页第一步即返回 HTTP 403，程序正确识别为 security_challenge 并执行退避重试。此时尚未进入 GraphQL、字段解析或 MySQL 写入。",
        PALE_YELLOW,
    )
    doc.add_page_break()

    doc.add_heading("1. 使用范围与处理原则", level=1)
    doc.add_paragraph(
        "本手册用于排查交付代码在目标机器上出现的 HTTP 403、安全验证、代理 407、网络超时和批量任务失败。"
        "重点是先确认请求实际走哪条网络，再判断是环境问题、代理问题、代码兼容问题还是 Kickstarter 页面与风控策略变化。"
    )
    for text in (
        "浏览器能打开页面，不等于 Python HTTP 请求能够被接受。浏览器会执行 JavaScript、保存验证 Cookie，并携带更完整的环境状态。",
        "不采用无限重试，不尝试自动绕过交互式验证码；遇到持续挑战应暂停并人工排查。",
        "先单 URL、后 3 URL、再 10 条小批量，最后才扩大任务规模。",
        "任何代理必须来自合法授权渠道，并遵守 Kickstarter 条款、访问限制及适用法律。",
    ):
        add_bullet(doc, text)

    doc.add_heading("2. 已知事实与当前缺口", level=1)
    add_table(
        doc,
        ["项目", "当前事实", "影响"],
        [
            ("现场结果", "3 个测试 URL 均在主页 GET 返回 403/security_challenge", "不是数据库或字段解析故障"),
            ("浏览器", "同机浏览器可以打开 Kickstarter", "浏览器可能已有 Cookie、JS 挑战状态或不同网络路径"),
            ("快速测试", "smoke_test.py 当前直接调用 collect(url)", "不会读取 config.json 中的代理配置"),
            ("批量程序", "metadata_batch_spider.py 会读取 config.json", "快速测试和批量程序可能走不同网络路径"),
            ("代理预检", "现有 proxy_probe.py 未与全部配置和批量启动门禁完全联动", "预检通过与实际任务配置可能不一致"),
            ("小象节点", "已知 Host/10030 可连，但无认证测试曾返回 407", "需要应用凭据或公网 IP 白名单"),
        ],
        [1500, 4200, 3660],
    )
    add_callout(
        doc,
        "最重要的结论",
        "不要通过修改数据库、URL 或字段解析来处理当前 403。先完成第 3 至第 7 章的网络与运行环境对照。",
        PALE_RED,
    )

    doc.add_heading("3. 第一次现场操作：停止重试并保存证据", level=1)
    add_step(doc, "停止连续运行", "如果三个项目全部 403，立即停止 smoke_test.py 和其他 Kickstarter 脚本，至少暂停 30 至 60 分钟。不要同时开启多个脚本。")
    add_step(doc, "保存输出", "保留完整终端日志、运行时间、URL、HTTP 状态、最终跳转地址和错误标签。不得在截图中暴露代理密码或数据库密码。")
    add_step(doc, "保存结果文件", "保留 smoke_test_results.json，并另存为带日期时间的文件，便于与更换网络或代理后的结果对比。")
    add_step(doc, "记录现场条件", "填写下表，尤其是网络类型、VPN、系统代理、公司网关和杀毒软件 HTTPS 检查。")
    add_table(
        doc,
        ["检查项", "需要记录的内容", "示例"],
        [
            ("机器", "Windows 版本、主机用途", "Windows 11 / 办公网电脑"),
            ("Python", "版本、虚拟环境路径", "Python 3.11 / .venv"),
            ("依赖", "curl_cffi 版本", "0.15.0"),
            ("网络", "家庭宽带、校园网、公司网、云服务器或手机热点", "公司网"),
            ("出口", "直连公网 IP、代理公网 IP", "分别记录，不公开凭据"),
            ("中间设备", "VPN、加速器、代理插件、HTTPS 检查、防火墙", "有/无及产品名称"),
            ("时间", "首次 403 和最后一次重试时间", "精确到分钟"),
        ],
        [1500, 4860, 3000],
    )

    doc.add_heading("4. 第二次现场操作：核对 Python 与依赖", level=1)
    doc.add_paragraph("在项目目录打开 PowerShell，执行以下命令：")
    add_code(
        doc,
        [
            r".\.venv\Scripts\python.exe -c \"import sys,curl_cffi; print(sys.version); print('curl_cffi',curl_cffi.__version__)\"",
            r".\.venv\Scripts\python.exe -m pip show curl_cffi lxml pymysql beautifulsoup4 requests tls-client",
            r".\.venv\Scripts\python.exe -m pip check",
        ],
    )
    for text in (
        "如果 .venv 指向另一台机器的 Python 路径，应删除旧虚拟环境后在本机重新创建；不要直接复制其他机器的 .venv。",
        "如果缺少依赖，重新执行 requirements-metadata.txt 安装；安装结束后必须执行 pip check。",
        "如果 Python 或 curl_cffi 版本与交付测试环境明显不同，先记录差异，不要一次性升级所有依赖。",
        "确认 Windows 日期、时间和时区正确；时间严重偏差可能导致 TLS 或认证异常。",
    ):
        add_bullet(doc, text)

    doc.add_heading("5. 第三次现场操作：确认实际网络路径", level=1)
    doc.add_heading("5.1 检查环境变量和 WinHTTP 代理", level=2)
    add_code(
        doc,
        [
            "Get-ChildItem Env: | Where-Object { $_.Name -match 'PROXY' }",
            "netsh winhttp show proxy",
        ],
    )
    doc.add_paragraph(
        "还应检查 Windows“设置 - 网络和 Internet - 代理”、浏览器代理扩展、VPN/加速器以及 IDE 自身的代理设置。"
        "浏览器、PowerShell、Python 和 Docker 可能各自使用不同代理，因此必须分别核对。"
    )
    doc.add_heading("5.2 比较浏览器与 Python 的公网出口", level=2)
    for text in (
        "在浏览器中打开可信的公网 IP 查询页面并记录地址。",
        "使用 Python 或现有代理预检程序查询公网出口并记录地址。",
        "若两者不同，先找出哪个代理、VPN或网关造成分流；在网络路径未统一前，不能用浏览器结果证明脚本直连可用。",
    ):
        add_step(doc, "执行", text)
    add_callout(
        doc,
        "注意",
        "公网 IP 属于运行诊断信息，可在内部工单记录，但不应与代理密码、数据库密码一起公开。",
        LIGHT_BLUE,
    )

    doc.add_heading("6. 第四次现场操作：单 URL、单请求对照", level=1)
    doc.add_paragraph("暂停期结束后，只测试一个项目，不要直接重新运行三个 URL：")
    add_code(
        doc,
        [
            r".\.venv\Scripts\python.exe project_metadata_spider.py --url \"https://www.kickstarter.com/projects/uchibacoyapiece/rosalie\"",
        ],
    )
    add_table(
        doc,
        ["结果", "判断", "下一步"],
        [
            ("主页 200，GraphQL 200", "直连当前可用", "再运行 3 URL 冒烟测试"),
            ("主页 403", "IP、网络、Cookie/JS 状态或请求指纹被拦截", "执行第 7 章网络切换对照"),
            ("主页 200，GraphQL 403", "接口策略、CSRF、会话或频率问题", "保留主页和 GraphQL 分步日志，检查页面结构与请求头"),
            ("超时/DNS 失败", "网络、防火墙、DNS 或代理污染", "先修复基础网络，不进入解析调试"),
            ("200 但 project_data_not_found", "页面结构变化或返回了未识别的验证页面", "保存响应标题与脱敏 HTML，更新解析规则"),
        ],
        [1900, 3660, 3800],
    )

    doc.add_heading("7. 第五次现场操作：更换网络进行 A/B 对照", level=1)
    add_step(doc, "保持代码不变", "不要同时修改 User-Agent、依赖、代理和 URL，否则无法判断哪个因素有效。")
    add_step(doc, "切换网络", "优先用手机热点或另一条正常家庭宽带；退出公司 VPN、加速器和浏览器代理扩展后再测试。")
    add_step(doc, "仍只测一个 URL", "使用第 6 章同一条命令和同一 URL。")
    add_step(doc, "记录 A/B 结果", "记录原网络和新网络的公网出口、状态码、响应时间和错误标签。")
    add_table(
        doc,
        ["原网络", "新网络", "较可能结论"],
        [
            ("403", "200", "原出口 IP、公司网关、地区线路或 HTTPS 检查导致"),
            ("403", "403", "脚本请求特征、依赖兼容、目标策略变化或两条线路均被限制"),
            ("超时", "200", "原网络 DNS、防火墙或路由故障"),
            ("200", "403", "新网络出口信誉或地区可用性更差"),
        ],
        [1700, 1700, 5960],
    )

    doc.add_heading("8. 第六次现场操作：合法代理配置路线", level=1)
    doc.add_paragraph(
        "只有直连持续不可用且业务确需代理时，才进入本路线。代理不能替代低频访问、失败退避和人工监控。"
    )
    doc.add_heading("8.1 必须取得的信息", level=2)
    for text in (
        "代理协议：HTTP 代理是否支持 HTTPS CONNECT。",
        "代理 Host 和端口。",
        "认证方式：应用 ID/用户名 + 应用密码，或运行机器公网 IP 白名单。",
        "产品是否有效、剩余流量/请求数、允许地区、并发和会话保持规则。",
        "是否支持同一项目的主页和 GraphQL 请求保持同一出口 IP。",
    ):
        add_bullet(doc, text)
    add_callout(
        doc,
        "现有小象代理信息",
        "端口为 10030，已知四个动态 Host。仅有 Host 和端口不够；无认证曾返回 407，需要取得应用凭据或完成公网 IP 白名单。",
        PALE_YELLOW,
    )
    doc.add_heading("8.2 填写 config.json", level=2)
    add_code(
        doc,
        [
            '"proxy_config": {',
            '  "proxy_host_list": ["http-dynamic.xiaoxiangdaili.com"],',
            '  "proxy_port": 10030,',
            '  "proxy_username": "实际应用ID或用户名",',
            '  "proxy_pwd": "实际应用密码",',
            '  "enable_proxy": 1',
            "}",
        ],
    )
    doc.add_paragraph(
        "config.json 只保存在运行机器，不进入公开仓库或交付截图。若使用 IP 白名单，应确认填写的是公网出口 IP，不是 127.0.0.1、192.168.x.x 或 10.x.x.x。"
    )
    doc.add_heading("8.3 代理预检和显式单项目测试", level=2)
    add_code(
        doc,
        [
            r".\.venv\Scripts\python.exe proxy_probe.py",
            r".\.venv\Scripts\python.exe project_metadata_spider.py --url \"https://www.kickstarter.com/projects/uchibacoyapiece/rosalie\" --proxy \"http://用户名:密码@代理Host:10030\"",
        ],
    )
    add_callout(
        doc,
        "程序差异",
        "当前 smoke_test.py 不读取 config.json 代理配置，因此不能直接用于验证配置文件里的代理是否生效。当前 proxy_probe.py 也未与所有配置及批量启动门禁完全联动。正式完善时应统一由同一配置生成代理 URL。",
        PALE_RED,
    )
    doc.add_heading("8.4 代理通过后的小批量验证", level=2)
    add_code(
        doc,
        [
            r".\.venv\Scripts\python.exe metadata_batch_spider.py --limit 3",
            r".\.venv\Scripts\python.exe metadata_batch_spider.py --limit 10",
        ],
    )
    doc.add_paragraph("只有 3 条和 10 条任务均稳定、字段完整且无连续挑战，才能逐步扩大到 50 条计划批次。")

    doc.add_heading("9. 推荐的软件完善项", level=1)
    add_table(
        doc,
        ["优先级", "完善项", "验收标准"],
        [
            ("P0", "smoke_test.py 增加 --config/--proxy，并读取与批量程序相同的代理配置", "快速测试和批量程序显示同一出口 IP"),
            ("P0", "proxy_probe.py 读取 config.json、支持认证并返回明确退出码", "DNS/TCP/认证/HTTPS/目标页分层输出"),
            ("P0", "批量启动前调用预检；失败时直接阻断任务", "预检非零退出码时不领取数据库任务"),
            ("P1", "诊断日志增加依赖版本、代理是否启用、脱敏 Host、出口 IP 摘要", "日志可复现且不泄露密码"),
            ("P1", "可选保存 403 响应摘要、标题、Server/Retry-After 等信息", "能区分代理 407、目标 403 和验证页"),
            ("P1", "支持关闭单次请求内部重试的诊断模式", "A/B 测试不会产生多次访问"),
            ("P2", "浏览器 Cookie 导入仅作为人工诊断选项并注明时效与合规边界", "不自动破解验证码，不长期保存敏感 Cookie"),
        ],
        [900, 5100, 3360],
    )

    doc.add_heading("10. 恢复运行的完整顺序", level=1)
    for title, detail in (
        ("环境检查", "Python、依赖、系统时间、网络代理均已核对。"),
        ("基础网络", "DNS、HTTPS 和公网出口可确认，无未知代理污染。"),
        ("单 URL", "主页与 GraphQL 均成功，核心字段可解析。"),
        ("3 URL 冒烟", "3/3 passed，退出码为 0。"),
        ("数据库初始化", "配置正确，schema 已初始化，能够写入测试项目。"),
        ("URL 文件 dry-run", "合法 URL 数量、规范化和去重结果符合预期。"),
        ("3 条批量", "任务状态、重试和数据库写入正常。"),
        ("10 条批量", "无连续 403/429，失败数可解释。"),
        ("计划批次", "逐步提高至 50 条，继续保持单线程和项目间随机等待。"),
    ):
        add_step(doc, title, detail)
    add_callout(
        doc,
        "启动门槛",
        "任一阶段出现环境级 403、407 或持续 429，就回到对应排查步骤，不允许跳过失败直接扩大批次。",
        PALE_RED,
    )

    doc.add_page_break()
    doc.add_heading("11. 后续可能发生的问题及预测", level=1)
    doc.add_paragraph(
        "以下问题按实际运行中最可能出现的层级排列。预测不代表一定发生，但应提前准备监控、日志和恢复方案。"
    )
    add_table(
        doc,
        ["风险", "可能表现", "原因", "建议措施"],
        [
            ("再次 403/挑战", "前几条成功，随后整批失败", "IP信誉、会话变化、频率或策略更新", "暂停、延长退避、检查出口和响应摘要"),
            ("HTTP 429", "返回 Too Many Requests", "访问频率或短时连接数过高", "遵循 Retry-After，降低批次并延长间隔"),
            ("代理 407", "所有代理请求立即失败", "凭据错误、白名单不匹配、产品到期", "检查应用凭据、公网 IP 和套餐"),
            ("代理 502/503/超时", "部分节点不稳定", "代理上游或线路故障", "有限重试、切换备用 Host、记录节点质量"),
            ("代理 IP 频繁切换", "主页成功、GraphQL 失败", "相关请求不在同一会话/出口", "使用项目级粘性会话"),
            ("浏览器正常而脚本失败", "人工页面 200、程序 403", "Cookie/JS/网络路径不同", "重新执行本手册第 5 至第 7 章"),
            ("页面结构变化", "200 但 project_data_not_found", "current_project 位置或 HTML 结构改变", "保存脱敏样本并更新解析器"),
            ("GraphQL 变化", "errors、字段为空或返回结构不同", "查询字段、权限或 schema 改变", "更新查询并为缺失字段提供兼容处理"),
            ("项目删除/地区限制", "404、重定向或页面内容不同", "项目不可用、URL 失效或区域差异", "标记 not_found/blocked，不无限重试"),
            ("字段天然为空", "地区、简介、外链缺失", "创作者未公开资料", "保留 NULL/空列表，不推测补全"),
            ("协作者分类不可靠", "无法判断是否服务商或类别", "Kickstarter 无统一标准字段", "保存原始公开资料，后处理分类并记录置信度"),
            ("MySQL 连接失败", "Access denied、Unknown database", "配置、用户、密码或容器未启动", "核对 config、Compose 和数据库权限"),
            ("任务状态遗留", "processing 长时间不恢复", "进程异常退出", "使用超时恢复机制并检查任务锁"),
            ("重复或漏采", "相同 URL 多条或年份项目不全", "URL 来源、规范化或候选集不完整", "导入 dry-run、去重；明确 --year 不负责发现 URL"),
            ("磁盘/日志增长", "运行变慢或磁盘满", "结果、日志和数据库持续累积", "设置保留周期、轮转日志和备份"),
            ("依赖升级回归", "新机器行为不同", "curl_cffi/Python/解析库版本变化", "锁定版本，升级前跑 3 URL 回归"),
            ("隐私与凭据泄露", "配置或日志出现密码", "截图、压缩包或仓库误提交", "忽略 config.json、脱敏日志、定期轮换凭据"),
        ],
        [1450, 2250, 2700, 2960],
    )

    doc.add_heading("12. 错误码快速判定矩阵", level=1)
    add_table(
        doc,
        ["现象", "请求到达位置", "优先检查", "是否继续重试"],
        [
            ("DNS 失败", "未到代理/目标", "Host、DNS、网络", "否，先修网络"),
            ("连接拒绝/超时", "可能未到代理", "端口、防火墙、代理线路", "有限，先换备用线路"),
            ("407", "到达代理，未通过认证", "凭据、白名单、套餐", "否，修复认证"),
            ("403/security_challenge", "通常到达 Kickstarter", "出口 IP、会话、TLS/HTTP 指纹、频率", "否，暂停并对照"),
            ("429", "到达 Kickstarter", "请求频率、Retry-After", "按服务端等待时间"),
            ("502/503", "代理上游或目标临时故障", "节点状态、目标状态", "有限退避"),
            ("200 但字段为空", "已收到页面", "是否验证页、页面结构、解析规则", "否，先检查内容"),
        ],
        [1900, 2200, 2760, 2500],
    )

    doc.add_heading("13. MySQL 与批量任务的后续问题", level=1)
    for text in (
        "config.example.json、config.json 和 docker-compose.metadata.yml 的数据库名、用户与密码必须保持一致；正式环境建议使用最小权限专用用户。",
        "首次运行前初始化 metadata_schema.sql；表结构升级时必须先备份。",
        "批量程序出现 failed > 0 时应返回非零退出码，自动任务不能只看是否生成日志。",
        "年份筛选只过滤已经导入的候选 URL，不会自动发现某一年全部 Kickstarter 项目。",
        "filtered_year 任务在切换年份或取消年份筛选后应恢复处理，不能永久标记 success。",
        "字段范围调整后，应同步更新解析结果、数据库表结构、写入 SQL、验证字段和文档。",
    ):
        add_bullet(doc, text)

    doc.add_heading("14. 运行监控建议", level=1)
    add_table(
        doc,
        ["指标", "建议观察值", "触发处理"],
        [
            ("HTTP 200 比例", "按批次统计主页与 GraphQL", "连续多个项目下降时暂停"),
            ("403/429 比例", "单独统计，不与普通解析错误混合", "出现环境级连续失败即停止"),
            ("代理 407", "应为 0", "任一出现即检查认证"),
            ("平均请求时间", "记录基线和趋势", "明显升高时检查线路"),
            ("任务状态", "pending/processing/retry/blocked/not_found", "processing 超时或 blocked 增加时人工介入"),
            ("字段缺失率", "按项目、创作者、协作者字段统计", "结构性升高时检查页面/API变化"),
            ("数据新鲜度", "保存采集时间", "筹款结束后补采最终金额"),
            ("代理消耗", "流量、请求数、余额、到期日", "接近阈值时告警，不自动无限续费"),
        ],
        [1900, 3900, 3560],
    )

    doc.add_heading("15. 最终验收清单", level=1)
    checklist = (
        "目标机器已重新创建虚拟环境，依赖安装和 pip check 正常。",
        "浏览器、PowerShell、Python 和批量程序的网络路径已经说明并记录。",
        "直连或代理的实际公网出口已确认。",
        "代理启用时，认证、HTTPS CONNECT 和目标页面预检全部通过。",
        "单项目主页与 GraphQL 均为 200，项目、创作者和协作者结构可解析。",
        "smoke_test.py 的 3 个项目全部通过，退出码为 0。",
        "URL 文件 dry-run、规范化和去重正确。",
        "3 条、10 条批量任务完成，数据库写入和任务状态符合预期。",
        "出现失败时退出码非零，日志不泄露凭据。",
        "已明确 URL 全量来源、年份筛选和协作者分类不属于同一个处理阶段。",
        "已建立暂停、重试、blocked 人工处理和数据备份流程。",
    )
    for item in checklist:
        add_bullet(doc, "□ " + item)

    doc.add_heading("16. 对外沟通模板", level=1)
    add_callout(
        doc,
        "可直接回复",
        "目前日志表明程序已经正常识别 Kickstarter 返回的 HTTP 403 安全验证，失败发生在项目主页第一次请求，尚未进入 GraphQL、字段解析或数据库写入。浏览器可以访问并不能证明 Python 请求一定可用，因为两者的 Cookie、JavaScript、TLS/HTTP 指纹和网络代理路径可能不同。请先暂停连续重试，按手册记录 Python/curl_cffi 版本、系统与环境变量代理、浏览器和 Python 的公网出口，并使用手机热点对同一单项目做 A/B 测试。若需使用代理，还必须取得有效应用凭据或完成公网 IP 白名单；仅有 Host 和端口会返回 407。代理预检和单项目测试通过后，再依次执行 3 条、10 条和计划批次。",
        PALE_GREEN,
    )

    doc.add_heading("17. 交付边界", level=1)
    doc.add_paragraph(
        "程序可以识别挑战、限制重试、记录失败并恢复任务，但无法保证 Kickstarter 在所有机器、网络和时间永久返回 200，"
        "也不应自动破解交互式验证码。历史 3/3 通过结果证明的是交付测试时的代码与网络组合可用，不是对未来网络环境的永久承诺。"
        "当目标站页面、GraphQL 或风控规则变化时，需要更新请求配置、解析规则或运行策略。"
    )

    core = doc.core_properties
    core.title = "Kickstarter 爬虫 HTTP 403 排查操作步骤与后续风险预测手册"
    core.subject = "浏览器可访问但 Python 爬虫返回 403 的操作型排查指南"
    core.author = "技术交付"
    core.keywords = "Kickstarter, HTTP 403, security_challenge, 代理, smoke_test, 排查"
    core.comments = "不包含代理密码、数据库密码或其他运行凭据"

    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
