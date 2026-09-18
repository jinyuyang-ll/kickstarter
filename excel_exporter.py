"""Export a work order's URL and metadata results to one Excel workbook."""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from metadata_batch_spider import connect_db, load_config


def csv_rows(path):
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        return list(csv.DictReader(source))


def append_table(book, title, rows):
    sheet = book.create_sheet(title[:31])
    if not rows:
        sheet.append(["说明"])
        sheet.append(["暂无数据"])
        return sheet
    columns = list(rows[0])
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3157D5")
    for row in rows:
        sheet.append([serialize_cell(row.get(column)) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(columns, 1):
        width = max(len(str(column)), *(len(str(serialize_cell(row.get(column)) or "")) for row in rows[:200]))
        sheet.column_dimensions[get_column_letter(index)].width = min(60, max(10, width + 2))
    return sheet


def serialize_cell(value):
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return value


def query_rows(db, sql, values=()):
    with db.cursor() as cursor:
        cursor.execute(sql, values)
        return list(cursor.fetchall())


def metadata_for_urls(config_path, urls):
    if not urls:
        return {}, "没有URL，未查询元数据"
    try:
        db = connect_db(load_config(config_path))
    except Exception as exc:
        return {}, "数据库不可用，工作簿仅包含URL和批次：" + str(exc)
    try:
        projects = []
        normalized = [url.split("?", 1)[0] for url in urls]
        for start in range(0, len(normalized), 500):
            chunk = normalized[start:start + 500]
            placeholders = ",".join(["%s"] * len(chunk))
            projects.extend(query_rows(db, f"""
                SELECT project_id, source_pid, project_link, title, currency,
                       goal_amount, pledged_amount, funding_start_date,
                       funding_end_date, backers_count, project_status,
                       creator_id, first_crawled_at, updated_at
                FROM project_metadata
                WHERE project_link IN ({placeholders})
                ORDER BY project_id
            """, chunk))
        project_ids = [row["project_id"] for row in projects]
        creator_ids = sorted({row["creator_id"] for row in projects if row.get("creator_id")})
        creators = []
        for start in range(0, len(creator_ids), 500):
            chunk = creator_ids[start:start + 500]
            placeholders = ",".join(["%s"] * len(chunk))
            creators.extend(query_rows(db, f"""
                SELECT creator_id, creator_name, profile_url, joined_date,
                       created_project_count, backed_project_count,
                       total_backers_across_projects, location, biography,
                       external_links, verified_identity, updated_at
                FROM creator_metadata WHERE creator_id IN ({placeholders})
            """, chunk))
        collaborators = []
        if project_ids:
            for start in range(0, len(project_ids), 500):
                chunk = project_ids[start:start + 500]
                placeholders = ",".join(["%s"] * len(chunk))
                collaborators.extend(query_rows(db, f"""
                    SELECT pc.project_id, pc.collaborator_id,
                           c.collaborator_name, c.profile_url,
                           pc.collaborator_role, c.location,
                           c.created_project_count, c.backed_project_count,
                           c.total_backers_across_projects,
                           c.is_service_provider, c.service_category,
                           c.classification_confidence
                    FROM project_collaborator pc
                    JOIN collaborator_metadata c
                      ON c.collaborator_id=pc.collaborator_id
                    WHERE pc.project_id IN ({placeholders})
                """, chunk))
        return {"项目数据": projects, "创作者": creators, "协作者": collaborators}, None
    except Exception as exc:
        return {}, "数据库表不可用，工作簿仅包含URL和批次：" + str(exc)
    finally:
        db.close()


def export_workbook(csv_path, output_path, config_path="config.json", manifest_path=None):
    csv_path = Path(csv_path).resolve()
    output_path = Path(output_path).resolve()
    rows = csv_rows(csv_path)
    urls = [row.get("project_url") for row in rows if row.get("project_url")]
    if manifest_path is None:
        manifest_path = csv_path.parent / (csv_path.stem + ".batches") / "manifest.json"
    else:
        manifest_path = Path(manifest_path)
    batches = []
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        for item in manifest.get("batches", []):
            batches.append({
                "batch_id": item.get("id"),
                "label": item.get("label"),
                "selected": item.get("selected"),
                "executed": item.get("executed"),
                "complete": item.get("complete"),
                "next_page": item.get("next_page"),
                "total_urls": item.get("total_urls"),
                "total_hits": item.get("total_hits"),
                "stop_reason": item.get("stop_reason"),
                "stop_label": item.get("stop_label"),
                "error": item.get("error"),
            })
    metadata, warning = metadata_for_urls(config_path, urls)
    book = Workbook()
    book.remove(book.active)
    append_table(book, "URL清单", rows)
    append_table(book, "批次状态", batches)
    for title, data in metadata.items():
        append_table(book, title, data)
    info = [{"项目": "生成时间", "值": __import__("datetime").datetime.now().isoformat(timespec="seconds")},
            {"项目": "URL数量", "值": len(urls)},
            {"项目": "数据库说明", "值": warning or "已包含当前URL对应的数据库元数据"}]
    append_table(book, "说明", info)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(output_path.suffix + ".tmp")
    book.save(temporary)
    temporary.replace(output_path)
    return {"output": str(output_path), "url_count": len(urls),
            "project_count": len(metadata.get("项目数据", [])), "warning": warning}


def main():
    parser = argparse.ArgumentParser(description="Export URL and metadata results to Excel")
    parser.add_argument("csv")
    parser.add_argument("--output", required=True)
    parser.add_argument("--config", default="config.json")
    parser.add_argument("--manifest")
    args = parser.parse_args()
    result = export_workbook(args.csv, args.output, args.config, args.manifest)
    print("EXCEL_SUMMARY " + json.dumps(result, ensure_ascii=False), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
