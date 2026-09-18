"""Regression tests for adaptive plans, persistent work orders and Excel output."""
import csv
import json
import tempfile
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook

import adaptive_planner as adaptive
import excel_exporter
import pipeline_runner
import project_metadata_spider as project_spider
import url_collector as collector
import web_app
from work_order_queue import QueueWorker, WorkOrderStore


class AdaptivePlannerTests(unittest.TestCase):
    def test_oversized_batch_is_replaced_by_collectable_children(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with patch.object(web_app, "ROOT", root):
                plan = web_app.make_plan({
                    "output": "urls.csv",
                    "states": ["successful"],
                    "request_mode": "http",
                    "max_pages": 1,
                    "delay_min": 8,
                    "delay_max": 8,
                })
            plan_path = root / "plan.json"
            collector.write_json_atomic(plan_path, plan)
            session = Mock()
            transport = Mock(proxy_url=None, secrets=[])
            transport.acquire.return_value = session

            def response(_session, url, *_args):
                query = parse_qs(urlparse(url).query)
                split = "goal_min" in query or "goal_max" in query
                return {"projects": [], "has_more": False, "total_hits": 1000 if split else 5000}

            with patch("adaptive_planner.collector.request_page", side_effect=response):
                result = adaptive.analyze_plan(plan_path, safe_pages=180, progress=lambda _: None, transport=transport)

            updated = json.loads(plan_path.read_text(encoding="utf-8"))
            self.assertEqual(result["split_events"], 1)
            self.assertEqual(result["leaf_batches"], 2)
            self.assertEqual(updated["adaptive_status"], "ready")
            self.assertEqual(len(updated["selected_batch_ids"]), 2)
            self.assertTrue(all(item["preflight"]["estimated_pages"] <= 180 for item in updated["batches"]))
            checkpoints = list((root / "urls.batches").glob("*.checkpoint.json"))
            self.assertEqual(len(checkpoints), 2)
            self.assertTrue(all(json.loads(p.read_text(encoding="utf-8"))["pending"]["total_hits"] == 1000 for p in checkpoints))
            with patch("url_collector.request_page", side_effect=AssertionError("预检首页不得重复请求")):
                collected = collector.collect_plan(plan_path, lambda _: None)
            self.assertTrue(collected["selected_complete"])

    def test_bounded_range_splits_at_midpoint_and_preserves_other_filters(self):
        url = collector.BASE_URL + "?category_id=12&state%5B%5D=successful&goal_min=100&goal_max=500&page=1"
        field, cut, left, right = adaptive.choose_split(url)
        self.assertEqual((field, cut), ("goal", 300))
        self.assertEqual(parse_qs(urlparse(left).query)["category_id"], ["12"])
        self.assertEqual(parse_qs(urlparse(right).query)["state[]"], ["successful"])
        self.assertEqual(parse_qs(urlparse(left).query)["goal_max"], ["300"])
        self.assertEqual(parse_qs(urlparse(right).query)["goal_min"], ["300"])


class WorkOrderTests(unittest.TestCase):
    def test_multiple_orders_are_persistent_and_claimed_serially(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "orders.sqlite3"
            store = WorkOrderStore(path)
            first = store.create("one", {}, "one.json", repeat_hours=24)
            second = store.create("two", {}, "two.json")
            claimed = store.claim_due()
            self.assertEqual(claimed["id"], first["id"])
            self.assertIsNone(store.claim_due())
            finished = store.finish(first["id"], 3)
            self.assertEqual(finished["status"], "queued")
            self.assertGreater(finished["next_run_at"], time.time())
            self.assertEqual(store.claim_due()["id"], second["id"])
            recovered = WorkOrderStore(path)
            self.assertEqual(recovered.get(second["id"])["status"], "queued")

    def test_worker_records_summary_and_reschedules_partial_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "pipeline_runner.py").write_text(
                "import json; print('PIPELINE_SUMMARY '+json.dumps({'status':'partial'})); raise SystemExit(3)",
                encoding="utf-8",
            )
            store = WorkOrderStore(root / "orders.sqlite3")
            created = store.create("one", {}, "plan.json", repeat_hours=1)
            claimed = store.claim_due()
            QueueWorker(store, root).run(claimed)
            result = store.get(created["id"])
            self.assertEqual(result["status"], "queued")
            self.assertEqual(result["last_exit_code"], 3)
            self.assertEqual(result["summary"]["status"], "partial")
            self.assertIn("WORK_ORDER_FINISHED", Path(result["log_path"]).read_text(encoding="utf-8"))

    def test_actions_keep_finished_order_reusable(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = WorkOrderStore(Path(tmp) / "orders.sqlite3")
            order = store.create("one", {}, "one.json", next_run_at=time.time() + 3600)
            paused = store.action(order["id"], "pause")
            self.assertEqual(paused["status"], "paused")
            resumed = store.action(order["id"], "run_now")
            self.assertEqual(resumed["status"], "queued")
            self.assertLessEqual(resumed["next_run_at"], time.time())


class ExcelTests(unittest.TestCase):
    def test_workbook_contains_urls_batches_and_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "urls.csv"
            with source.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=collector.CSV_FIELDS)
                writer.writeheader()
                writer.writerow({"project_url": "https://www.kickstarter.com/projects/u/p", "project_id": "1"})
            batch_dir = root / "urls.batches"
            batch_dir.mkdir()
            (batch_dir / "manifest.json").write_text(json.dumps({"batches": [{"id": "a", "label": "G1", "complete": True}]}), encoding="utf-8")
            output = root / "result.xlsx"
            fake = {"项目数据": [{"project_id": 1, "title": "Demo"}], "创作者": [], "协作者": []}
            with patch("excel_exporter.metadata_for_urls", return_value=(fake, None)):
                result = excel_exporter.export_workbook(source, output)
            self.assertEqual(result["url_count"], 1)
            book = load_workbook(output, read_only=True)
            self.assertEqual(set(book.sheetnames), {"URL清单", "批次状态", "项目数据", "创作者", "协作者", "说明"})
            self.assertEqual(book["URL清单"].max_row, 2)
            book.close()


class PipelineTests(unittest.TestCase):
    def test_partial_url_collection_exports_excel_but_does_not_start_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            plan_path = root / "plan.json"
            plan_path.write_text(json.dumps({"output": str(root / "urls.csv"), "batches": []}), encoding="utf-8")
            args = SimpleNamespace(
                plan_file=str(plan_path), auto_split=False, safe_pages=180,
                run_metadata=True, metadata_limit=0, metadata_year=None,
                excel_output=str(root / "result.xlsx"), config="config.json",
                schema="metadata_schema.sql",
            )
            collected = {"total_urls": 5, "selected_complete": False, "failed": False}
            with patch("pipeline_runner.url_collector.collect_plan", return_value=collected),                  patch("pipeline_runner.excel_exporter.export_workbook", return_value={"output": str(root / "result.xlsx"), "url_count": 5}),                  patch("pipeline_runner.url_importer.import_urls") as importer:
                code = pipeline_runner.run(args)
            self.assertEqual(code, 3)
            importer.assert_not_called()

    def test_security_challenge_is_not_retried_by_metadata_request(self):
        session = Mock()
        response = Mock(status_code=403, text="<html>captcha</html>", url="https://www.kickstarter.com/")
        session.request.return_value = response
        with patch("project_metadata_spider.time.sleep") as sleep, self.assertRaisesRegex(RuntimeError, "security_challenge"):
            project_spider.request_with_retry(session, "get", response.url)
        self.assertEqual(session.request.call_count, 1)
        sleep.assert_not_called()


class WorkOrderWebTests(unittest.TestCase):
    def test_create_work_order_persists_plan_schedule_and_pipeline_options(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            store = WorkOrderStore(root / ".collection" / "orders.sqlite3")
            data = {
                "output": "urls.csv", "states": ["successful"],
                "request_mode": "http", "repeat_hours": 24,
                "auto_split": True, "safe_pages": 180,
                "run_metadata": True, "metadata_limit": 25,
                "excel_output": "result.xlsx",
                "schedule_at": "2099-01-01T12:30",
            }
            with patch.object(web_app, "ROOT", root), patch.object(web_app, "ORDER_STORE", store):
                order = web_app.create_work_order(data)
            self.assertEqual(order["status"], "queued")
            self.assertEqual(order["repeat_hours"], 24)
            self.assertTrue(order["spec"]["auto_split"])
            self.assertTrue(order["spec"]["run_metadata"])
            self.assertEqual(order["spec"]["metadata_limit"], 25)
            self.assertEqual(len(order["batches"]), 1)
            self.assertTrue(Path(order["plan_path"]).exists())


if __name__ == "__main__":
    unittest.main()
