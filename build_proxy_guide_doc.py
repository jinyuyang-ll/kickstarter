from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from build_delivery_doc import (
    BLUE,
    DARK_BLUE,
    GRAY,
    LIGHT_BLUE,
    NAVY,
    add_code_block,
    add_page_number,
    bullet,
    number,
    set_repeat_table_header,
    set_table_geometry,
    shade_cell,
    set_font,
    style_document,
)


OUTPUT = "HTTP代理通用注册购买配置与任务前预检指南.docx"


def build():
    doc = Document()
    style_document(doc)
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    set_font(header.add_run("HTTP代理通用注册、购买、配置与预检指南"), size=9, bold=True, color=GRAY)
    add_page_number(section.footer.paragraphs[0])

    for _ in range(3):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(kicker.add_run("运行人员操作手册"), size=11, bold=True, color=BLUE)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    set_font(title.add_run("HTTP代理通用配置与任务前预检"), size=27, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(45)
    set_font(subtitle.add_run("供应商选择、购买、认证、预检、迁移与故障排查"), size=14, color=DARK_BLUE)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(meta.add_run("资料核对日期：2026-07-28"), size=10.5, color=GRAY)
    doc.add_page_break()

    doc.add_heading("1. 先看结论", level=1)
    doc.add_paragraph(
        "当前四个动态代理Host和10030端口均可连接，但无认证测试返回407。407表示代理"
        "认证未通过，请求通常尚未到达Kickstarter。使用前必须完成应用ID/应用密码认证，"
        "或把运行机器的公网出口IP加入产品白名单。"
    )
    for text in (
        "不使用代理：config.json保持 enable_proxy=0，先运行直连冒烟测试。",
        "使用代理：先运行proxy_probe.py，预检通过后再启动批量任务。",
        "预检失败：停止本轮任务，检查套餐、凭据、白名单、Host和端口。",
    ):
        bullet(doc, text)

    doc.add_heading("2. 注册与购买", level=1)
    for text in (
        "打开 https://www.xiaoxiangdaili.com，使用手机号注册并登录。",
        "按会员中心要求完成实名认证；官方帮助中心说明交易需要实名认证。",
        "进入产品购买页，选择“隧道代理—动态转发”。",
        "先申请试用或短周期体验，验证Kickstarter连通性后再购买正式套餐。",
        "确认有效期、计费方式、并发、IP切换和线路后下单支付。",
        "企业如需充值、对公转账、合同或发票，联系官方在线客服确认当前规则。",
        "付款后进入“产品管理—隧道代理”查看应用ID、应用密码和白名单入口。",
    ):
        number(doc, text)
    doc.add_paragraph(
        "平台套餐和优惠会变化，本指南不写死价格。最终金额、退款、合同和发票条件以购买"
        "页面、订单页面及客服确认为准。"
    )

    doc.add_heading("3. 认证方式", level=1)
    doc.add_heading("3.1 应用ID和应用密码", level=2)
    doc.add_paragraph(
        "后台产品管理页显示的应用ID和应用密码用于Proxy-Authorization。应用ID不是网站"
        "登录手机号，应用密码也不是网站登录密码。"
    )
    add_code_block(
        doc,
        [
            '"proxy_host_list": ["http-dynamic.xiaoxiangdaili.com"],',
            '"proxy_port": 10030,',
            '"proxy_username": "应用ID",',
            '"proxy_pwd": "应用密码",',
            '"enable_proxy": 1',
        ],
    )
    doc.add_heading("3.2 公网IP白名单", level=2)
    doc.add_paragraph(
        "把运行爬虫机器的公网出口IP绑定到产品白名单后，可不传用户名密码。不能填写"
        "127.0.0.1、192.168.x.x或10.x.x.x等本地/局域网地址。公网IP变化后需要更新"
        "白名单；公网IP经常变化时更适合账号密码认证。"
    )

    doc.add_heading("4. 动态转发节点", level=1)
    for text in (
        "http-dynamic.xiaoxiangdaili.com:10030",
        "http-dynamic-s02.xiaoxiangdaili.com:10030",
        "http-dynamic-s03.xiaoxiangdaili.com:10030",
        "http-dynamic-s04.xiaoxiangdaili.com:10030",
    ):
        bullet(doc, text)
    doc.add_paragraph(
        "同一个项目的主页与GraphQL请求应尽量保持同一代理Host、同一出口和同一Session，"
        "避免在短时间内频繁切换身份。"
    )

    doc.add_heading("5. 任务前代理预检", level=1)
    doc.add_paragraph(
        "交付包已有proxy_probe.py，可作为独立启动门禁。本指南只规定运行流程，不修改"
        "爬虫代码。"
    )
    for text in (
        "检查配置是否启用代理，Host、端口和认证字段是否完整。",
        "检查DNS解析和10030端口。",
        "通过代理查询公网IP，确认代理真正生效。",
        "通过代理访问Kickstarter测试项目，确认不是403、407或验证页。",
        "全部通过后才启动metadata_batch_spider.py。",
    ):
        number(doc, text)
    add_code_block(
        doc,
        [
            r".\.venv\Scripts\python.exe proxy_probe.py",
            "if ($LASTEXITCODE -ne 0) {",
            '    throw "代理预检失败，已停止批量任务"',
            "}",
            r".\.venv\Scripts\python.exe metadata_batch_spider.py --limit 10",
        ],
    )
    doc.add_paragraph(
        "如果配置明确关闭代理，应先运行smoke_test.py作为直连门禁；3个测试项目全部通过"
        "后再启动批量任务。"
    )

    doc.add_heading("6. 407排查", level=1)
    for text in (
        "确认产品未过期且购买的是隧道动态转发。",
        "确认使用的是应用ID和应用密码，而非网站登录密码。",
        "白名单认证时，确认绑定的是当前运行机器的公网出口IP。",
        "确认Host拼写和10030端口正确。",
        "确认代理配置没有被系统环境变量或其他配置文件覆盖。",
        "重新运行proxy_probe.py并保存错误结果。",
        "仍为407时，向官方客服提供应用、检查时间、Host和错误信息。",
    ):
        number(doc, text)

    doc.add_heading("7. 采购与安全注意事项", level=1)
    for text in (
        "先试用和小批量验证，再购买长期套餐。",
        "确认代理服务用途符合平台条款、目标网站条款和适用法律。",
        "不要把config.json、应用密码或数据库密码提交到公开仓库。",
        "不要通过公开聊天发送长期有效凭据；凭据泄露后立即重置。",
        "代理可用不等于目标网站一定返回200，仍需保留限速、重试和人工暂停机制。",
    ):
        bullet(doc, text)

    doc.add_heading("8. 官方资料", level=1)
    for text in (
        "快速入门：https://www.xiaoxiangdaili.com/help/dev/index",
        "隧道接入指南：https://www.xiaoxiangdaili.com/tunnel-short/CSharp/guide",
        "新用户问题：https://www.xiaoxiangdaili.com/help/question/index",
        "购买问题：https://www.xiaoxiangdaili.com/help/question/common-buy",
        "407排查：https://www.xiaoxiangdaili.com/help/4274",
    ):
        bullet(doc, text)

    doc.add_page_break()
    doc.add_heading("9. 代理产品基础知识", level=1)
    doc.add_paragraph(
        "不同供应商的产品名称并不统一。采购前不要只看“HTTP代理”四个字，应确认代理协议、"
        "出口网络、IP生命周期、会话保持、计费方式和认证方式。"
    )
    type_table = doc.add_table(rows=1, cols=4)
    type_table.style = "Table Grid"
    for i, value in enumerate(("类型", "特点", "适用情况", "注意事项")):
        shade_cell(type_table.rows[0].cells[i], LIGHT_BLUE)
        set_font(type_table.rows[0].cells[i].paragraphs[0].add_run(value), size=9.5, bold=True, color=NAVY)
    rows = (
        ("HTTP代理", "转发HTTP请求", "普通网页和接口", "HTTPS目标通常仍需支持CONNECT"),
        ("HTTPS隧道", "通过CONNECT建立加密隧道", "HTTPS网站", "代理能看到目标域名和连接元数据"),
        ("SOCKS5", "工作在更低层，协议适用范围广", "非HTTP客户端或特殊网络", "现有程序需确认客户端是否支持"),
        ("静态/固定IP", "出口长期不变", "白名单、登录态、长期会话", "单一IP风险集中"),
        ("动态/轮换IP", "出口按请求、时间或会话变化", "分布式读取和地域测试", "切换过快可能破坏Cookie和会话"),
        ("数据中心IP", "速度快、成本低", "低敏感目标和一般接口", "某些网站识别率较高"),
        ("住宅/ISP IP", "更接近普通用户网络", "需要真实地域和较高兼容性", "通常价格更高，必须核查来源合规性"),
    )
    for values in rows:
        cells = type_table.add_row().cells
        for i, value in enumerate(values):
            set_font(cells[i].paragraphs[0].add_run(value), size=8.5)
    set_repeat_table_header(type_table.rows[0])
    set_table_geometry(type_table, [1500, 2450, 2450, 2960])

    doc.add_heading("10. 供应商选择清单", level=1)
    doc.add_paragraph("至少对比以下项目，不建议仅按最低价格购买：")
    for text in (
        "是否明确支持HTTP CONNECT访问HTTPS目标，是否提供测试地址和错误码文档；",
        "出口国家/地区、运营商、数据中心或住宅属性是否满足业务需要；",
        "是否支持粘性会话，粘性时长能否覆盖一个项目的主页与GraphQL请求；",
        "认证方式是否支持账号密码、IP白名单或两者同时支持；",
        "计费单位是流量、请求数、IP数、端口数、并发数还是使用时长；",
        "并发、速率、每秒新连接数和单IP请求限制；",
        "是否提供使用量、失败率、出口IP和账单API；",
        "服务可用性、技术支持响应、退款、试用、合同、发票和对公付款；",
        "隐私政策、日志保留、IP来源、授权链路及允许用途；",
        "是否允许目标站点所在地区访问，是否存在地区出口或内容限制。",
    ):
        bullet(doc, text)
    doc.add_paragraph(
        "建议先用3至10个真实项目做小规模验收，记录成功率、延迟、407/403比例和实际消耗，"
        "再决定套餐和供应商。"
    )

    doc.add_heading("11. 通用采购与开通流程", level=1)
    for text in (
        "确认目标站点、预计URL数量、单项目请求数、运行周期和允许地区。",
        "列出必须能力：HTTPS CONNECT、粘性会话、认证方式、并发和出口地区。",
        "选择2至3家供应商申请试用，不要直接购买长期大套餐。",
        "完成账号注册、实名认证或企业认证；保存合同主体和客服渠道。",
        "购买最小可验证规格，并确认试用流量、有效期和自动续费设置。",
        "从产品后台取得Host、端口、用户名/应用ID、密码/密钥和白名单入口。",
        "在非生产机器完成代理预检和3个项目冒烟测试。",
        "核对账单消耗与预估是否一致，再扩大批次。",
        "形成内部交接记录：账号归属、续费日期、套餐限制、凭据保管人和应急联系人。",
    ):
        number(doc, text)

    doc.add_heading("12. 通用配置字段映射", level=1)
    mapping = doc.add_table(rows=1, cols=3)
    mapping.style = "Table Grid"
    for i, value in enumerate(("程序字段", "供应商后台可能名称", "说明")):
        shade_cell(mapping.rows[0].cells[i], LIGHT_BLUE)
        set_font(mapping.rows[0].cells[i].paragraphs[0].add_run(value), size=9.5, bold=True, color=NAVY)
    for values in (
        ("proxy_host_list", "Gateway / Endpoint / Host / 隧道地址", "不要包含路径；注意主备线路"),
        ("proxy_port", "Port / 端口", "HTTP和SOCKS端口可能不同"),
        ("proxy_username", "Username / App ID / Tunnel ID", "通常不是网站登录账号"),
        ("proxy_pwd", "Password / App Secret / Key", "不要写入公开仓库"),
        ("enable_proxy", "本地开关", "只有预检通过后才设为1"),
        ("地区参数", "Country / Region / Zone", "部分供应商编码在用户名中"),
        ("会话参数", "Session / Sticky ID", "部分供应商编码在用户名或Host中"),
    ):
        cells = mapping.add_row().cells
        for i, value in enumerate(values):
            set_font(cells[i].paragraphs[0].add_run(value), size=8.5)
    set_repeat_table_header(mapping.rows[0])
    set_table_geometry(mapping, [2200, 3200, 3960])
    doc.add_paragraph(
        "有些供应商把国家、城市、会话ID和生命周期编码进用户名，例如"
        "customer-zone-country-session。不能机械照抄其他供应商格式，应按本供应商文档组合。"
    )

    doc.add_heading("13. 认证方式与适用选择", level=1)
    for text in (
        "账号密码：适合公网IP经常变化、云函数、多机器或无法维护白名单的环境。",
        "IP白名单：适合固定公网出口、凭据不便下发的服务器；公网IP变化后必须更新。",
        "令牌/API Key：部分供应商使用请求头或用户名字段携带令牌，需确认客户端支持方式。",
        "混合模式：不同机器可分别使用白名单和账号密码，但要确认产品是否允许。",
    ):
        bullet(doc, text)
    doc.add_paragraph(
        "登录后台账号密码与代理认证凭据通常不是同一组。407优先检查代理凭据和白名单，"
        "不要反复修改网站登录密码。"
    )

    doc.add_heading("14. 会话保持与IP轮换策略", level=1)
    doc.add_paragraph(
        "代理并非切换越快越好。对于当前采集流程，一个项目通常包含主页和GraphQL两个相关"
        "请求，应尽量在同一会话、同一Cookie和同一出口IP内完成。"
    )
    for text in (
        "项目级粘性：同一项目的全部请求使用同一个代理会话，项目结束后再考虑切换。",
        "批次级粘性：一小批项目复用同一出口，适合稳定性较好的出口和低请求频率。",
        "失败切换：连接超时、502/503等网络故障可换线路；403时先暂停，不要立刻高速换IP重试。",
        "认证失败不切换：407是凭据或白名单问题，换大量出口通常无效。",
        "不要混用Cookie：出口IP切换后继续共享不同线程的Cookie可能形成异常会话特征。",
        "遵循网站限制：代理不能替代限速、失败退避、robots和目标网站条款。",
    ):
        bullet(doc, text)

    doc.add_heading("15. 分层代理预检", level=1)
    check_table = doc.add_table(rows=1, cols=4)
    check_table.style = "Table Grid"
    for i, value in enumerate(("层级", "检查内容", "通过标准", "失败处理")):
        shade_cell(check_table.rows[0].cells[i], LIGHT_BLUE)
        set_font(check_table.rows[0].cells[i].paragraphs[0].add_run(value), size=9.5, bold=True, color=NAVY)
    checks = (
        ("L0 配置", "开关、Host、端口、认证字段", "格式完整且符合认证模式", "不启动任务"),
        ("L1 网络", "DNS、TCP端口", "解析成功且端口可连接", "查防火墙、DNS、套餐"),
        ("L2 认证", "代理出口IP查询", "返回200且获得出口IP", "查407、白名单、凭据"),
        ("L3 HTTPS", "CONNECT和证书", "HTTPS目标正常建立", "查协议、证书、代理能力"),
        ("L4 目标", "Kickstarter测试页", "返回200且不是验证页", "暂停并分析403/429"),
        ("L5 数据", "3个项目冒烟测试", "字段完整、失败率可接受", "不进入全量"),
        ("L6 成本", "消耗和限额", "符合预算、无异常扣费", "调整套餐或频率"),
    )
    for values in checks:
        cells = check_table.add_row().cells
        for i, value in enumerate(values):
            set_font(cells[i].paragraphs[0].add_run(value), size=8.2)
    set_repeat_table_header(check_table.rows[0])
    set_table_geometry(check_table, [1450, 2600, 2650, 2660])
    doc.add_paragraph(
        "自动化任务的启动门禁建议至少要求L0至L4通过；正式扩大批量前要求L5和L6通过。"
    )

    doc.add_heading("16. 常见错误与定位", level=1)
    error_table = doc.add_table(rows=1, cols=3)
    error_table.style = "Table Grid"
    for i, value in enumerate(("现象", "常见原因", "优先处理")):
        shade_cell(error_table.rows[0].cells[i], LIGHT_BLUE)
        set_font(error_table.rows[0].cells[i].paragraphs[0].add_run(value), size=9.5, bold=True, color=NAVY)
    errors = (
        ("DNS解析失败", "Host错误、本地DNS或域名故障", "核对控制台Host；更换DNS前先查供应商状态"),
        ("连接被拒绝", "端口错误、产品未开通、防火墙", "核对端口和套餐；测试TCP"),
        ("连接超时", "线路不通、出口拥堵、防火墙", "换官方备用线路；检查本机网络"),
        ("407", "凭据错误、白名单不匹配、产品过期", "核对应用凭据/公网IP和套餐"),
        ("400/405", "代理协议或请求格式不兼容", "核对HTTP、SOCKS和CONNECT用法"),
        ("403", "目标站点拒绝、验证页、地区或IP信誉", "停止高频重试；保留响应并换低频测试"),
        ("429", "请求速率超过限制", "延长间隔；遵循Retry-After"),
        ("502/503", "代理上游或目标暂时不可用", "有限重试；换备用线路"),
        ("TLS/证书错误", "中间人证书、协议版本、系统时间", "不要关闭证书校验；查代理说明和系统时间"),
        ("出口IP未变化", "代理未生效或被环境变量覆盖", "查询直连/代理IP；检查实际传参"),
        ("200但字段为空", "返回验证页、地区页或HTML结构变化", "检查页面标题和关键标记，不只看状态码"),
    )
    for values in errors:
        cells = error_table.add_row().cells
        for i, value in enumerate(values):
            set_font(cells[i].paragraphs[0].add_run(value), size=8.2)
    set_repeat_table_header(error_table.rows[0])
    set_table_geometry(error_table, [1800, 3500, 4060])

    doc.add_heading("17. 更换代理供应商时的迁移步骤", level=1)
    for text in (
        "保留原配置备份，但删除或吊销旧凭据，记录旧服务终止时间。",
        "取得新供应商Host、端口、协议、认证和地区/会话参数说明。",
        "只修改代理配置映射，不同时修改采集字段、数据库和线程策略。",
        "依次完成L0至L6预检，并与旧供应商比较成功率、延迟和成本。",
        "使用相同的3个测试URL做A/B验证，避免测试样本不同造成误判。",
        "确认新供应商稳定后再扩大批量；保留短期回退方案。",
        "更新内部凭据保管、续费提醒、客服联系人和故障记录。",
    ):
        number(doc, text)

    doc.add_heading("18. 成本与容量估算", level=1)
    doc.add_paragraph(
        "采购前先估算请求量。当前新元数据流程在年份符合时通常每个项目2次核心请求；"
        "年份不符合时仅请求主页1次。评论、更新、奖励等旧脚本会产生额外分页请求，不能"
        "按2次估算。"
    )
    for text in (
        "按流量计费：用小批量实测平均响应字节，乘以项目数并预留重试和日志开销。",
        "按请求计费：项目数乘以主页、GraphQL及分页请求数，再预留失败重试比例。",
        "按并发计费：单线程运行不代表无新连接，应确认连接复用和每秒新连接限制。",
        "按IP/端口/时长计费：确认过期后处理、自动续费和剩余资源是否结转。",
        "设置预算告警、日用量上限和异常消耗检查，避免错误循环快速消耗套餐。",
    ):
        bullet(doc, text)

    doc.add_heading("19. 安全、合规与凭据管理", level=1)
    for text in (
        "只使用来源、授权和用途明确的代理资源；核查供应商隐私与日志政策。",
        "遵守Kickstarter服务条款、robots、访问限制和适用法律。",
        "代理凭据使用最小权限；测试、生产和个人环境分开。",
        "config.json只保存在运行机器，不进入压缩包或公共仓库。",
        "凭据定期轮换；人员离职、设备丢失或疑似泄露时立即吊销。",
        "日志中隐藏完整密码、令牌和带认证信息的代理URL。",
        "不要通过关闭TLS证书校验来绕过证书问题。",
        "保留采购、授权、变更、故障和数据用途记录，便于审计。",
    ):
        bullet(doc, text)

    doc.add_heading("20. 上线验收清单", level=1)
    for text in (
        "□ 代理产品、有效期、计费和自动续费已确认。",
        "□ Host、端口、协议、应用凭据或公网IP白名单已确认。",
        "□ 直连IP与代理出口IP已分别记录。",
        "□ DNS、端口、认证、HTTPS和Kickstarter目标测试均通过。",
        "□ 3个真实项目冒烟测试全部通过。",
        "□ 单线程、请求间隔、失败退避和最大重试符合要求。",
        "□ 日志不会泄露密码或认证URL。",
        "□ 预算上限、用量告警、续费日期和负责人已登记。",
        "□ 代理预检失败时，批量任务不会由操作人员继续启动。",
        "□ 已准备供应商客服渠道和切换/回退方案。",
    ):
        bullet(doc, text)

    doc.add_heading("21. 文档适用边界", level=1)
    doc.add_paragraph(
        "本指南是供应商无关的采购、配置和排查框架，小象代理仅作为现有项目实例。其他"
        "供应商的字段名称、认证拼接、会话参数和购买流程可能不同，应以其官方文档为准。"
        "代理预检通过只能证明当前网络和认证可用，不能保证目标网站永久返回200，也不能"
        "替代限速、失败退避和人工监控。"
    )

    doc.core_properties.title = "HTTP代理通用注册购买配置与任务前预检指南"
    doc.core_properties.subject = "代理选型、采购、认证、预检、迁移和故障排查"
    doc.core_properties.author = "项目交付"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
